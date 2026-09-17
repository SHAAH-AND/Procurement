from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\MK\Documents\Procurement")
OUT = ROOT / "output" / "documents" / "Galle_Face_Hotel_ProcureFlow_SRS_Revision_and_Development_Requirements.docx"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F4F7FB"
GRAY = "D9D9D9"
MIDGRAY = "666666"
RED = "9C0006"
AMBER = "9C6500"
GREEN = "006100"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRAY, size="5"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_keep_with_next(paragraph, keep=True):
    paragraph.paragraph_format.keep_with_next = keep


def set_repeat_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_text(p, text, bold=False, italic=False, color=BLACK, size=None):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return r


def add_bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25 + (0.2 * level))
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    add_text(p, "•  ")
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_number(doc, text, number, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28 + (0.2 * level))
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    add_text(p, f"{number}.  ")
    add_text(p, text)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def paragraph(doc, text="", bold_lead=None, italic=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):], italic=italic)
    else:
        add_text(p, text, italic=italic)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        set_cell_shading(c, NAVY)
        set_cell_border(c)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, bold=True, color=WHITE, size=font_size)
        if widths:
            c.width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        set_repeat_no_split(table.rows[-1])
        for i, value in enumerate(row):
            c = cells[i]
            c.text = ""
            if ri % 2:
                set_cell_shading(c, PALE)
            set_cell_border(c)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if i == 0 and len(str(value)) < 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, str(value), size=font_size)
            if widths:
                c.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_table(doc, rows):
    return add_table(doc, ["Area", "Delivery status", "Current capability", "Required work"], rows,
                     widths=[1.25, 1.2, 2.15, 2.1], font_size=8.0)


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 10.5)):
    s = styles[name]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 15)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True

for sty in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
    styles[sty].font.name = "Aptos"
    styles[sty]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles[sty]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles[sty].font.size = Pt(9.5)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(18)
title_ppr = p._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
add_text(p, "Galle Face Hotel ProcureFlow SRS Revision and Development Requirements", bold=True, size=25)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(30)
add_text(p, "Instructions for correcting the Software Requirements Specification and completing the system build", size=13, color=MIDGRAY)

meta = add_table(doc, ["Document item", "Value"], [
    ["Prepared for", "The author and reviewers of the Galle Face Hotel ProcureFlow SRS"],
    ["Source SRS", "Software Requirement Specification for Galle Face Hotel Version 1.0 dated 16 September 2026"],
    ["System reviewed", "ProcureFlow application built on Zoho Catalyst"],
    ["Review date", "17 September 2026"],
    ["Document purpose", "Revision instruction, development scope and acceptance baseline"],
    ["Recommended status", "Use as the change brief for SRS Version 1.1"],
], widths=[1.7, 4.95], font_size=9)

doc.add_paragraph()
p = doc.add_paragraph()
add_text(p, "Main conclusion", bold=True, size=10.5)
p.paragraph_format.space_after = Pt(4)
paragraph(doc, "The requested business process can be implemented on Zoho Catalyst using the existing ProcureFlow application as the baseline. The present SRS and the present application are not yet aligned closely enough to claim exact implementation. The SRS must first resolve the conflicts and omissions listed in this document. Development must then close the functional, security, data, integration and operational gaps and demonstrate them through traceable acceptance tests.")

add_page_break(doc)

heading(doc, "Contents", 1)
for item in [
    "1 Purpose and required outcome",
    "2 Decisions required before the SRS is revised",
    "3 Mandatory SRS corrections",
    "4 Required functional specification by module",
    "5 Required nonfunctional specification",
    "6 Complete development scope",
    "7 Required data model and business rules",
    "8 Integrations and external dependencies",
    "9 Migration configuration and operational readiness",
    "10 Verification acceptance and traceability",
    "11 Definition of complete delivery",
    "12 Instructions to the SRS author",
    "Appendix A Current ProcureFlow baseline and confirmed gaps",
    "Appendix B Recommended requirement wording patterns",
]:
    add_bullet(doc, item)

heading(doc, "Status labels used in this document", 2)
add_table(doc, ["Label", "Meaning"], [
    ["Supported baseline", "The current application contains a meaningful implementation, but it must still be validated against the final requirement."],
    ["Enhancement required", "The current implementation must be extended or corrected."],
    ["New development", "The requirement is not implemented as an end-to-end feature."],
    ["Decision required", "The SRS contains conflicting or incomplete instructions and must be clarified before build."],
    ["External dependency", "Delivery depends on a third-party service, credential, commercial agreement or client decision."],
], widths=[1.4, 5.25], font_size=9)

heading(doc, "1 Purpose and required outcome", 1)
paragraph(doc, "This document is written for the person revising the Galle Face Hotel ProcureFlow SRS. It identifies what must change in the specification and describes the complete development scope needed to deliver the agreed system without known functional gaps.")
paragraph(doc, "The revised SRS should become the contractual and testable baseline for design, development, configuration, migration, security review, user acceptance testing and production sign-off. Every final requirement should state the triggering condition, permitted actor, required data, validation rule, resulting status, audit event, error behavior and acceptance test.")
paragraph(doc, "The SRS author should not describe a capability as complete merely because a screen, setting or upload control exists. The requirement is complete only when the server enforces the rule, the user interface explains it, the event is audited, failure behavior is defined and an acceptance test proves the result.")

heading(doc, "Recommended target architecture", 2)
paragraph(doc, "Use one dedicated Catalyst project for the Galle Face Hotel group unless Cloud Partners formally decides to operate one shared multi-customer ProcureFlow deployment. Within the hotel-group project, property and legal-entity access must remain server-scoped. This recommendation matches the current deployment design and reduces cross-customer isolation risk. If the shared SaaS model is selected instead, the SRS must add membership, organization selection, tenant suspension, tenant-level configuration, tenant-aware support access, metering and cross-tenant security tests as mandatory scope.")

heading(doc, "2 Decisions required before the SRS is revised", 1)
paragraph(doc, "The following decisions change the data model, workflow engine or integration boundary. They must be approved by Galle Face Hotel before the revised SRS is signed.")
decisions = [
    ["D01", "Deployment model", "The SRS describes both a shared multi-tenant SaaS platform and a dedicated Galle Face Hotel tenant. The current checkout uses one installation per hotel group.", "Adopt one dedicated Catalyst project for the group and remove shared-SaaS behavior from this client SRS. Keep property and legal-entity segregation inside the project."],
    ["D02", "Approval authority", "The budget-based route in sections 3.1.3.5 and 5 differs from the CAPEX OPEX and repair route shown later. The document also says seven approvers while listing six actions.", "Approve one authoritative workflow matrix by expenditure type, budget condition, legal entity and value threshold. State whether each stage approves, reviews, negotiates or authorizes."],
    ["D03", "Direct purchase order", "The scope calls RFQ optional, while the award section says a PO can only be generated after an RFQ reaches Awarded.", "Allow direct PO only through a documented exception route such as emergency, sole source or contract call-off. Require RFQ for all other purchases."],
    ["D04", "Expenditure categories", "The SRS calls the list three fixed values but lists Capex, Opex, Repair and AMC. Later attachment rules introduce Services.", "Use five controlled values: CapEx, OpEx Goods, Services, Repairs and AMC. If OpEx must include goods and services, define Services as a subcategory and revise every rule consistently."],
    ["D05", "Budget conditions", "The PR screen lists Budgeted and Non-Budgeted, but approval logic also uses Exceeds Budget.", "Use three system-calculated results: Within Budget, Not Budgeted and Exceeds Budget. The requester may select a proposed budget line, but may not choose the final result."],
    ["D06", "Bid confidentiality", "The SRS says bids are sealed but also allows procurement to view bids and negotiate while the RFQ remains open.", "Define the opening model. Recommended: bid contents stay hidden until the deadline; authorized procurement opens them afterward; a formal negotiation round may then reopen only shortlisted bids with full version history."],
    ["D07", "Receipt completion", "The SRS excludes rejected quantities from accepted stock and matching but also says a PO can become Fulfilled when quantities are fully accepted or accounted for.", "Use Partially Received, Received With Exception and Fulfilled. Fulfilled requires accepted quantity equal to ordered quantity or an approved short-close decision."],
    ["D08", "Payment boundary", "The SRS refers to automated payments and EFT bank data, but also describes recording payments.", "Confirm whether Release 1 only records externally executed payments. Treat actual bank payment initiation as a separate integration unless a named bank API and control process are approved."],
    ["D09", "WhatsApp delivery", "The SRS requires WhatsApp invitation links without naming an approved provider or consent policy.", "Select an approved WhatsApp Business provider, message-template process, sender, consent basis and delivery-failure handling. Otherwise label WhatsApp as a later-phase dependency."],
    ["D10", "Accounting integration", "Zoho Books is described as an example or optional integration.", "State whether Books is mandatory for first go-live, which records are system-of-record, the sync direction, field mapping and failure reconciliation process."],
]
add_table(doc, ["ID", "Decision", "Problem to resolve", "Recommended SRS position"], decisions,
          widths=[0.5, 1.25, 2.35, 2.6], font_size=7.8)

add_page_break(doc)
heading(doc, "3 Mandatory SRS corrections", 1)
paragraph(doc, "The SRS author should apply every correction below before requesting technical sign-off. Priority P0 means the current wording can produce a materially different system. Priority P1 means the omission prevents objective acceptance. Priority P2 covers quality and editorial consistency.")

corrections = [
    ["SRS-01", "P0", "Document control", "Version 0.1, Version 1.0 and the date 2 August 2026 appear together. Owner, issue date and approvals are blank.", "Set one document ID, version, issue date, owner, review status and approval table. Record Version 1.1 changes in Document History."],
    ["SRS-02", "P2", "Names and terminology", "The document alternates among Procure-Flow, ProcureFlow, Galle Face Hotel, Galle Face Retreat and customer placeholders.", "Use ProcureFlow and Galle Face Hotel consistently. Replace all placeholders and incorrect entity names."],
    ["SRS-03", "P2", "Structure and numbering", "Section 3.1.3.11 is duplicated, Payment Module is numbered 3.1.1.10 and several IDs are missing or duplicated.", "Renumber sections and requirements. Assign a unique stable identifier to every SHALL statement."],
    ["SRS-04", "P1", "Priority model", "Section 1.4 says every requirement states its priority, but the requirements do not.", "Add Must, Should or Could to every requirement, or remove the claim and use a release field in the traceability matrix."],
    ["SRS-05", "P0", "Architecture", "Shared SaaS and dedicated installation language conflict.", "State the selected deployment and tenancy model, identity boundary, property scope and data isolation approach."],
    ["SRS-06", "P0", "Approval rules", "Different sections define different chains and responsibilities.", "Replace prose examples with one authoritative decision table and state how role vacancies, delegation, absence, rejection, recall and resubmission work."],
    ["SRS-07", "P0", "Budget determination", "The user appears able to choose budget status even though later sections require an automatic budget check.", "Specify server-calculated budget status using property, legal entity, department, financial year, period, expense type and budget line."],
    ["SRS-08", "P0", "RFQ and PO route", "RFQ is both optional and mandatory.", "Define standard and exception sourcing routes, approval evidence and the status transitions for each."],
    ["SRS-09", "P0", "Sealed bids", "Confidentiality, opening time, visibility and negotiation behavior are not defined.", "Add precise bid-visibility and versioning rules, including administrator access and audit events."],
    ["SRS-10", "P1", "Attachment enforcement", "Required document lists exist, but the rules do not define configurable slots, file checks, waivers or who may override them.", "Define a document-requirement matrix by stage and category, allowed formats and sizes, waiver authority, rejection behavior and audit requirements."],
    ["SRS-11", "P0", "Supplier onboarding", "Vendor registration, internal vendor access, vendor activation and supplier approval are mixed together.", "Define separate statuses for invited, registration in progress, pending SCM review, pending Finance review, active, rejected, suspended and blacklisted."],
    ["SRS-12", "P1", "Purchase receipt", "PO eligibility and fulfillment rules are internally inconsistent.", "Define the PO and GRN state machine, cumulative quantity rules, accepted and rejected handling, short close and reversal or correction rules."],
    ["SRS-13", "P0", "Invoice matching", "The SRS requires line-level quantity and amount matching but only identifies a few invoice header inputs in the functional table.", "Define invoice lines, taxes, freight, discounts, tolerances, duplicate detection, partial invoicing, multiple GRNs, services and authorized mismatch override."],
    ["SRS-14", "P0", "Payments and credits", "The effect of credits and discrepancy overrides is incomplete.", "Specify credit allocation, available balance, partial application, payment authorization, overpayment prevention, reversal and settlement status rules."],
    ["SRS-15", "P1", "Custom modules", "The SRS implies administrators can configure arbitrary workflows and modules without defining supported behavior.", "Limit Release 1 to a defined schema builder and permissions, or specify a full workflow engine including states, actions, validations and audit behavior."],
    ["SRS-16", "P1", "Reports", "Dashboards are described by name but report dimensions, filters, calculations, export format and data freshness are not specified.", "Add a report catalogue with calculation definitions, security scope, filters, drill-down, export and reconciliation criteria."],
    ["SRS-17", "P1", "Integrations", "Books, email, WhatsApp, EFT and webhooks lack complete interface contracts.", "For each integration, state direction, authentication, field mapping, trigger, idempotency, retry, error queue, audit and ownership."],
    ["SRS-18", "P1", "Security", "Security principles are stated, but measurable controls and test evidence are incomplete.", "Add session rules, MFA or SSO decision, access reviews, audit retention, file scanning, rate limiting, privacy, data residency, recovery and security-test requirements."],
    ["SRS-19", "P1", "Performance and availability", "No measurable service objectives are defined.", "Add response-time targets, expected concurrency and data volumes, availability target, maintenance approach, monitoring and capacity thresholds."],
    ["SRS-20", "P1", "Backup and recovery", "The SRS does not define recovery point or recovery time objectives.", "State backup frequency, retention, restore ownership, RPO, RTO and restore-test cadence."],
    ["SRS-21", "P1", "Migration", "Source files and open transactions are not defined in the SRS.", "Add source owners, templates, validation rules, duplicate handling, mock loads, reconciliation, cutover and rollback."],
    ["SRS-22", "P1", "Acceptance", "Most statements do not include observable pass conditions.", "Add Given When Then acceptance criteria and link every Must requirement to a test case and business owner."],
]
add_table(doc, ["ID", "Priority", "Topic", "Current problem", "Required revision"], corrections,
          widths=[0.55, 0.5, 1.15, 2.15, 2.3], font_size=7.4)

heading(doc, "Editorial corrections that affect meaning", 2)
for text in [
    "Change User Cause Diagram to User Case Diagram or Use Case Diagram and confirm every actor and system boundary shown in the diagram.",
    "Replace references to CRM records such as leads, contacts and deals in the operating environment section with procurement records.",
    "Correct the item currency default. The client context uses LKR, while the item section says USD. State the organization base currency and permitted transaction currencies.",
    "Correct the expenditure-value statement that says three values but lists four. Apply the approved category model throughout forms, attachments, workflows and reports.",
    "Replace undefined phrases such as internal vendor, external vendor, final department budget, automated payment process and dynamic field with defined terms.",
    "Correct the sentence stating that support activity can be attributed to the vendor's own support staff. Distinguish platform support engineers, client administrators and supplier users.",
    "Review minimum mobile and desktop versions at go-live. Old operating-system minimums should not substitute for a supported-browser matrix.",
    "Proofread spelling, grammar, spacing and capitalization after the business rules are corrected. Editorial cleanup must not hide unresolved logic conflicts.",
]:
    add_bullet(doc, text)

add_page_break(doc)
heading(doc, "4 Required functional specification by module", 1)
paragraph(doc, "This section defines the minimum content that the revised SRS must include for each module. The wording may be adapted, but the rules and acceptance boundaries should remain explicit.")

modules = [
    ("4.1 Identity users roles and property access", [
        "Use Catalyst hosted identity for internal users. Define invitation, first login, sign-out, password recovery and account deactivation behavior.",
        "Separate organizational role from permission profile. Define create, read, edit, delete, approve, configure, export and override permissions per module.",
        "Resolve property access on the server. Group users may see all properties; property users may see only assigned properties. Define cross-property approval exceptions.",
        "Define approver delegation, temporary absence, inactive users, vacant roles, duplicate role holders and escalation timers.",
        "State whether MFA, SSO and domain restrictions are required at go-live and identify the responsible identity administrator.",
    ]),
    ("4.2 Organization property legal entity and department setup", [
        "Model group, cluster, property, legal entity and department as separate controlled records with active dates and ownership.",
        "Map every property to one legal entity, delivery address, base currency, fiscal-year settings and repair authorization rule.",
        "Define who can edit organization profile, branding, tax identifiers, document headers and numbering sequences.",
        "Prevent deactivation of a property, department or legal entity while unresolved transactions depend on it, or define controlled archival behavior.",
    ]),
    ("4.3 Item and service master", [
        "Require Name, SKU and cost price and enforce tenant-unique SKU on the server.",
        "Support Goods and Service, purchasing UOM, conversion factor, lead time, tax treatment, preferred supplier, brand, description, specification, par level, warranty, perishable flag, maintenance eligibility and status.",
        "Use approved Department, Category and Subcategory cascades. Define whether one item may belong to multiple departments or properties.",
        "Prevent inactive or unapproved items from selection. Define how non-catalogue requests become approved items without losing request history.",
        "Add bulk import validation, row-level error reporting, duplicate handling and import audit logs.",
    ]),
    ("4.4 Purchase requisitions", [
        "Capture property, legal entity, department, expected date, expenditure category, delivery address, reference, notes and line details.",
        "Calculate each line as quantity multiplied by rate, less discount, plus tax. Recalculate totals on the server.",
        "For budgeted requests, select an approved budget line. For no-budget requests, require an approved reason and explanatory evidence.",
        "Render the approved CAPEX, OPEX and repair forms, including addition or replacement questions and entity-specific letterheads.",
        "Show relevant purchase history without exposing other properties or restricted supplier data.",
        "Define Draft, Pending Approval, Approved, Rejected, Recalled, Converted Partially, Converted Fully, Cancelled and Closed transitions.",
    ]),
    ("4.5 Approval engine", [
        "Store workflow definitions as versioned configuration rather than hard-coded route arrays.",
        "Select the route using legal entity, property, department, expenditure category, budget condition, value band and exception type.",
        "Persist the workflow version and complete route snapshot on each transaction so later configuration changes do not rewrite history.",
        "Require comments on rejection, override, delegation and bypass. Record actor, effective actor, action, timestamp, prior state and resulting state.",
        "Prevent self-approval where required and define segregation-of-duties rules for requester, approver, receiver, invoice matcher and payment recorder.",
    ]),
    ("4.6 Budget management", [
        "Support property, legal entity, department, financial year, expense type, active dates and period cadence.",
        "Implement CAPEX budget lines with reference, asset category A to J, item, addition or replacement, location, quantities, unit price, value, quarter, discretionary classification and business justification.",
        "Calculate approved, committed, actual, available and forecast amounts. Define exactly when commitment is reserved, released, moved to actual or reversed.",
        "Use concurrency protection so two simultaneous requests cannot both consume the same remaining budget.",
        "Define transfer, adjustment, closure, carry-forward and audit rules and require Finance approval where appropriate.",
    ]),
    ("4.7 RFQ and sourcing", [
        "Create RFQs in Draft, validate required fields and evidence, then publish to selected eligible suppliers.",
        "Require a deadline and timezone. Automatically close submissions when the deadline passes.",
        "Implement the approved sealed-bid model. Log bid opening, viewing, downloads and any deadline changes.",
        "Support supplier selection, bulk invitations, delivery requirements, evaluation criteria, clarifications and amendments.",
        "Provide bid comparison across price, tax, delivery, payment terms and technical scores. Preserve every submitted bid version.",
        "Require award evidence and approvals before one winning supplier is selected. Prevent a second award unless the original award is formally cancelled.",
    ]),
    ("4.8 Supplier registration and portal", [
        "Implement the complete registration form shown in the SRS, including entity, tax, contacts, categories, bank data, declarations and signature or stamp evidence.",
        "Apply conditional document requirements for company type, foreign currency and supplier category. Enforce file type and size on both client and server.",
        "Route submissions through SCM review and Finance review. Support requests for more information, resubmission, rejection and activation.",
        "Keep supplier authentication separate from employee access. Scope a supplier session to one supplier and one organization.",
        "Show only assigned RFQs, the supplier's own bids, issued POs, its invoices and its payment history. Never expose competitor bids.",
        "Provide secure single-use seven-day invitations, hashed credentials and sessions, revocation, rate limiting and security notifications.",
    ]),
    ("4.9 Purchase orders and supplier acknowledgement", [
        "Generate a PO from an awarded RFQ or an approved exception route. Copy the correct supplier, items, final prices, taxes, property, address, terms and delivery dates.",
        "Require approved decision evidence before issue. Store a rendered immutable issue version and retain later amendments separately.",
        "Allow the supplier to accept or reject. Rejection must notify procurement and move the PO to a defined resolution state.",
        "Define amendment, cancellation, reissue and alternate-award behavior and approval rules.",
    ]),
    ("4.10 Goods receipt and service completion", [
        "Allow receipts only against eligible issued or partly received POs.",
        "Require received quantity to equal accepted plus rejected quantity. Block cumulative receipt above ordered quantity unless an explicit tolerance and approval are configured.",
        "Support multiple partial receipts. Set Fulfilled only when accepted quantity reaches ordered quantity or an authorized short close is recorded.",
        "Require category-specific delivery, inspection, commissioning or service-completion documents before GRN or SES completion.",
        "Provide correction or reversal with approval and audit. Do not silently edit completed receipts.",
        "Create assets from accepted CapEx receipt lines when asset registration applies.",
    ]),
    ("4.11 Invoices three way match and credits", [
        "Capture all required invoice header and line fields, including supplier and purchaser tax data, delivery, place of supply, taxes and payment method.",
        "Validate the approved invoice-number and date formats. Prevent duplicates using supplier, invoice number and legal entity.",
        "Match invoice lines against PO price and cumulative accepted GRN or SES quantities. Support partial invoices and multiple receipts.",
        "Make price, quantity and tax tolerances configurable by category or legal entity. Show each discrepancy and its calculation.",
        "Block payment for unresolved discrepancies unless an authorized Finance override includes a reason and evidence.",
        "Support vendor credit issuance, remaining balance, application to one or more invoices, refund and cancellation without exceeding the available balance.",
    ]),
    ("4.12 Payments recurring bills and settlement", [
        "Support partial and full payment records, payment date, method, reference, voucher and settled invoices.",
        "Prevent cumulative payment above the invoice balance after credits and adjustments.",
        "Define approval and segregation of duties for recording, authorizing and reversing payments.",
        "Specify whether recurring bills create drafts, require approval and match to contract or PO before payment.",
        "If actual payment initiation is in scope, add beneficiary verification, dual authorization, bank response, reconciliation, failure and reversal requirements.",
    ]),
    ("4.13 Documents correspondence and templates", [
        "Use a configurable document-requirement matrix by transaction stage, expenditure category, legal entity and value band.",
        "Store files in secure Catalyst object storage under organization and record prefixes. Enforce authorization on every download.",
        "Define allowed extensions, MIME validation, file-size limits, virus-scanning decision, retention and deletion controls.",
        "Version PDF templates and correspondence templates. Persist the template version and rendered output issued to each recipient.",
        "Record delivery status for email and WhatsApp messages. Do not mark a transaction notified when delivery fails.",
    ]),
    ("4.14 Reporting audit and support", [
        "Provide named reports for purchasing, spend, budgets, payables aging, CapEx, OpEx, supplier performance, approvals and property or cluster performance.",
        "Define calculation sources and reconciliation totals for every financial report.",
        "Write append-only business audit events and protect them from ordinary edit and delete operations.",
        "Support constrained read-only support sessions only if the platform support model remains in scope. Require reason, 30-minute expiry, revocation, visible banner and client-visible audit entry.",
        "Provide export controls and audit sensitive exports, supplier bank access and administrative configuration changes.",
    ]),
]
for title, bullets in modules:
    heading(doc, title, 2)
    for b in bullets:
        add_bullet(doc, b)

add_page_break(doc)
heading(doc, "5 Required nonfunctional specification", 1)
paragraph(doc, "The revised SRS should replace general assurances with measurable targets. The values below are recommended baselines and must be confirmed against the purchased Catalyst plan, Galle Face Hotel policy and expected transaction volumes.")
nfr = [
    ["NFR-01", "Availability", "Define the monthly availability target and exclusions for planned maintenance. Recommended starting target: 99.9 percent for production application services."],
    ["NFR-02", "Performance", "Define response targets by operation. Recommended starting target: 95 percent of normal API requests within 2 seconds and standard dashboard loads within 4 seconds under agreed load."],
    ["NFR-03", "Capacity", "State named-user count, concurrent-user target, suppliers, items, annual PRs, RFQs, POs, invoices, attachments and retention horizon. Include a growth assumption."],
    ["NFR-04", "Scalability", "Require pagination and indexed queries for lists and exports. Do not assume a single query returns all Data Store rows; large exports must use controlled batching."],
    ["NFR-05", "Identity", "Define MFA or SSO, session duration, sign-out, inactive-user handling, invitation expiry and privileged-access review."],
    ["NFR-06", "Authorization", "Require server-side enforcement for every endpoint and object download, with property, legal-entity and supplier scoping."],
    ["NFR-07", "Data protection", "Require TLS in transit, Catalyst platform protection at rest, secret storage outside client code, least-privilege credentials and masked bank information in the interface and logs."],
    ["NFR-08", "Privacy", "Identify personal and banking data, permitted purposes, retention, export, correction and deletion responsibilities. Record applicable Sri Lankan privacy obligations and client policy."],
    ["NFR-09", "Audit", "Define mandatory event fields, append-only behavior, searchable retention period, export format, timezone handling and access rights."],
    ["NFR-10", "Backup and recovery", "Set backup cadence, retention, restore procedure, RPO and RTO. Require a documented restore test before go-live and periodically afterward."],
    ["NFR-11", "Resilience", "Define idempotency for create, award, issue, invoice, payment and integration operations. Specify retry and reconciliation behavior after partial failure."],
    ["NFR-12", "File security", "Validate size, extension and MIME type; authorize every upload and download; define malware scanning, quarantining and prohibited content."],
    ["NFR-13", "Accessibility", "Target WCAG 2.2 AA for keyboard access, focus visibility, labels, contrast, error messages, tables and responsive layouts."],
    ["NFR-14", "Compatibility", "State supported browsers and minimum versions at launch. Test responsive use on supported mobile browsers rather than relying only on operating-system names."],
    ["NFR-15", "Observability", "Define structured application logs, correlation IDs, error alerts, integration monitoring, audit review and health checks without exposing secrets."],
    ["NFR-16", "Secure development", "Require dependency review, secret scanning, static checks, authorization tests, tenant or property isolation tests, vulnerability remediation and release approval."],
    ["NFR-17", "Maintainability", "Require versioned schema, configuration export, deployment instructions, rollback, code review, automated tests and supported runtime versions."],
    ["NFR-18", "Localization", "Define timezone, date display, LKR formatting, tax labels, language and legal-entity document formats."],
    ["NFR-19", "Data quality", "Define required-field, uniqueness, reference-integrity, numeric-range and status-transition validation plus import reconciliation."],
    ["NFR-20", "Retention and archiving", "Define retention for transactions, bids, supplier documents, audit logs and support sessions. State when deletion is prohibited by an open transaction or audit requirement."],
]
add_table(doc, ["ID", "Area", "Requirement to add or confirm"], nfr, widths=[0.65, 1.35, 4.75], font_size=8.2)

heading(doc, "6 Complete development scope", 1)
paragraph(doc, "The following work packages cover the build required after the SRS decisions are approved. Each package includes backend enforcement, user interface behavior, audit events, automated tests, documentation and UAT evidence.")

work = [
    ["WP01", "Architecture and configuration", "Confirm dedicated or shared tenancy; define environments; version organization, property, legal-entity and module settings; remove obsolete architecture assumptions.", "P0"],
    ["WP02", "Identity and access control", "Complete internal invitations, role and profile permissions, property scope, delegation, inactive-user blocking, privileged access and optional SSO or MFA configuration.", "P0"],
    ["WP03", "Master data", "Complete properties, legal entities, departments, categories, items, suppliers, contacts, bank records, taxes, currencies, payment terms, numbering and imports.", "P0"],
    ["WP04", "Requisition forms", "Build CAPEX, OpEx, Service, Repair and AMC forms; implement conditional fields, purchase history, budget-line selection, evidence checks and exact PDF outputs.", "P0"],
    ["WP05", "Workflow engine", "Replace fixed routes with versioned rules by category, budget result, entity, property, department and amount; add delegation, vacancies, escalations and route snapshots.", "P0"],
    ["WP06", "Budget engine", "Build period and item-level budget lines, commitment ledger, actuals, release and reversal, concurrent-spend protection, transfers and consolidated reporting.", "P0"],
    ["WP07", "Document rules", "Create configurable requirement slots and server-side gates for PR, RFQ, bid, award, PO, receipt, invoice and supplier-registration stages.", "P0"],
    ["WP08", "RFQ and bid management", "Add draft and publish states, secure invitations, sealed bid behavior, deadlines, amendments, clarifications, negotiation rounds, comparisons, award evidence and notifications.", "P0"],
    ["WP09", "Supplier onboarding", "Build the full registration form, conditional evidence, SCM and Finance review, more-information cycle, activation and secure portal access.", "P0"],
    ["WP10", "Purchase orders", "Implement direct and RFQ-derived routes, immutable issued versions, amendments, supplier acknowledgement, rejection resolution and alternate award.", "P0"],
    ["WP11", "Receiving and assets", "Correct partial receipt behavior, implement exact quantity rules, required evidence, SES, short close, reversals and CapEx asset creation.", "P0"],
    ["WP12", "Invoices and matching", "Add invoice lines and complete tax fields, duplicate format validation, line-level three-way match, tolerances, partial billing, exception approval and rematch.", "P0"],
    ["WP13", "Payments and credits", "Enforce discrepancy block, credit application, outstanding balance, partial settlement, approvals, reversals and reconciliation. Add bank execution only if separately confirmed.", "P0"],
    ["WP14", "Documents and communications", "Complete secure object storage, permissions, template versioning, PDFs, email delivery, delivery logs and optional WhatsApp provider integration.", "P1"],
    ["WP15", "Reporting and analytics", "Implement the approved report catalogue, filters, drill-down, exports, calculations, data scope and financial reconciliation.", "P1"],
    ["WP16", "Audit and support", "Harden append-only audit events, sensitive-action logging, configuration history, export monitoring and constrained support sessions if retained.", "P0"],
    ["WP17", "Zoho Books and webhooks", "Complete field mapping, OAuth, sync ownership, idempotency, retries, error queue, manual replay, reconciliation and disconnect behavior.", "P1"],
    ["WP18", "Data migration", "Build templates, validations and controlled imports for properties, users, items, suppliers, budgets and approved opening commitments.", "P0"],
    ["WP19", "Quality and security", "Add unit, API, state-transition, authorization, property-isolation, supplier-scope, concurrency, integration, browser, performance and security tests.", "P0"],
    ["WP20", "Operations and adoption", "Prepare monitoring, backup and restore, release and rollback, administrator guide, role training, support process, pilot rollout and hypercare reporting.", "P0"],
]
add_table(doc, ["Work package", "Area", "Required development", "Priority"], work,
          widths=[0.75, 1.25, 4.25, 0.55], font_size=7.8)

heading(doc, "Recommended delivery sequence", 2)
for i, text in enumerate([
    "Approve the SRS decisions and workflow matrix before changing transactional code.",
    "Complete architecture, access control, master data and the versioned workflow and budget foundations.",
    "Complete the PR through award path, including required-document gates and supplier onboarding.",
    "Complete PO, receiving, invoice, matching, credits and payments with transaction-safe state changes.",
    "Complete reports, integrations, migration and operational controls.",
    "Run system integration testing, security testing and client UAT against the traceability matrix before production release.",
], 1):
    add_number(doc, text, i)

heading(doc, "7 Required data model and business rules", 1)
paragraph(doc, "The final schema may use different physical table names, but it must represent the following records and relationships without hiding material business data in unvalidated free-form JSON.")

data = [
    ["Organization structure", "Organization, cluster, property, legal entity, department, delivery location, financial year", "Effective dates, active status and ownership"],
    ["Access", "User, role, permission profile, property assignment, delegation, support session", "No authority derived from browser-supplied organization or property values"],
    ["Catalogue", "Item or service, SKU, category, subcategory, UOM, tax, supplier eligibility, compliance", "Unique SKU, active and approved state"],
    ["Supplier", "Supplier, contacts, categories, bank accounts, registration submission, review stages, documents", "Bank data access and changes separately audited"],
    ["Budget", "Budget header, period, CAPEX line, commitment ledger, adjustment, transfer", "Immutable ledger entries with reversals rather than silent balance edits"],
    ["Workflow", "Workflow definition, version, condition, stage, role, threshold, transaction route snapshot", "In-flight records keep their approved workflow version"],
    ["Requisition", "PR header, lines, category details, budget link, approvals, attachments", "Calculated totals and statuses enforced on server"],
    ["Sourcing", "RFQ, invited supplier, amendment, clarification, bid version, evaluation, award", "Competitor confidentiality and one active award"],
    ["Ordering", "PO, lines, issue version, amendment, supplier decision", "Issued documents remain reproducible"],
    ["Receiving", "GRN or SES, lines, inspection, accepted, rejected, short close, reversal", "Cumulative receipt and accepted quantity controls"],
    ["Payables", "Invoice, invoice lines, match result, discrepancy, override, credit, allocation, payment", "Outstanding balance calculated from auditable components"],
    ["Documents", "Attachment, document requirement, waiver, template version, rendered document, delivery event", "Authorization and retention follow parent record"],
    ["Operations", "Audit event, integration run, retry item, migration batch, reconciliation result", "Correlation IDs and immutable history"],
]
add_table(doc, ["Domain", "Required records", "Essential rule"], data, widths=[1.2, 3.15, 2.25], font_size=8)

heading(doc, "Transaction safety rules", 2)
for text in [
    "Award creation and PO generation must be idempotent. A retry may not create a second PO or second winning bid.",
    "Budget reservation must protect against two simultaneous submissions consuming the same funds. Use an atomic or compensating design approved for Catalyst Data Store.",
    "GRN creation must either save the header, all lines and the correct PO status or leave the transaction recoverable and visibly incomplete.",
    "Invoice creation, matching and budget settlement must be repeatable without double-counting actual spend.",
    "Payment and credit allocation must use current outstanding balances and prevent race-condition overpayment.",
    "Every integration create operation must carry an idempotency or external reference and keep a replayable failure record.",
]:
    add_bullet(doc, text)

heading(doc, "Required state machines", 2)
states = [
    ["Purchase request", "Draft, Pending Approval, Approved, Rejected, Recalled, Partially Converted, Fully Converted, Cancelled, Closed"],
    ["RFQ", "Draft, Published, Closed for Bids, Under Evaluation, Negotiation, Awarded, Cancelled"],
    ["Bid", "Draft, Submitted, Superseded, Withdrawn, Shortlisted, Awarded, Unsuccessful"],
    ["Supplier onboarding", "Invited, Registration in Progress, Pending SCM Review, More Information Required, Pending Finance Review, Approved, Rejected, Suspended, Blacklisted"],
    ["Purchase order", "Draft, Approved for Issue, Issued, Supplier Accepted, Supplier Rejected, Partially Received, Received With Exception, Fulfilled, Short Closed, Cancelled"],
    ["Invoice", "Draft, Pending Match, Matched, Review, Discrepancy, Override Approved, Partially Paid, Paid, Rejected, Cancelled"],
    ["Credit", "Draft, Open, Partially Applied, Fully Applied, Refunded, Cancelled"],
]
add_table(doc, ["Record", "Required states to define"], states, widths=[1.55, 5.05], font_size=8.5)

heading(doc, "8 Integrations and external dependencies", 1)
integrations = [
    ["Catalyst Authentication", "Internal identity and session", "Organization setup, invitation policy, ZAID per environment", "Client IT and implementation team"],
    ["Catalyst Data Store", "Transactional records", "Approved schema, indexes, query pagination, concurrency design", "Implementation team"],
    ["Catalyst Stratus", "Attachments and rendered documents", "Bucket, private access, path convention, retention, malware decision", "Implementation team and client security"],
    ["Email", "Invitations, RFQ, award, rejection, PO and alerts", "Verified sender, templates, recipient data, bounce handling", "Client IT and implementation team"],
    ["WhatsApp Business", "Optional supplier invitations and alerts", "Approved provider, sender, message templates, consent and commercial account", "Client and provider"],
    ["Zoho Books", "Optional vendor, item, bill, payment or reference sync", "Books organization, OAuth credentials, scopes, mapping and source-of-truth decision", "Client Finance and implementation team"],
    ["Bank or payment service", "Only if actual electronic payment is required", "Named provider, API, account, security approval, dual authorization and reconciliation", "Client Treasury and bank"],
    ["Webhooks", "Approved outbound business events", "Endpoint ownership, signing secret, retry and monitoring", "Receiving-system owner"],
    ["SSO", "Optional corporate sign-in", "Identity provider metadata, domain, group mapping and break-glass admin", "Client IT"],
]
add_table(doc, ["Dependency", "Purpose", "Required input or decision", "Owner"], integrations,
          widths=[1.25, 1.4, 2.9, 1.15], font_size=7.8)

heading(doc, "Integration acceptance rules", 2)
for text in [
    "No integration is accepted only because a connection test succeeds. UAT must prove correct field mapping, retry, duplicate prevention, failure visibility and reconciliation.",
    "Credentials must remain in Catalyst server-side configuration or approved connection storage and must never be returned to the browser or configuration export.",
    "The system must remain usable for core procurement when an optional integration is unavailable. Failed outbound work must remain visible and replayable.",
    "Every integration must have a named business owner, technical owner, support route and data-protection approval.",
]:
    add_bullet(doc, text)

heading(doc, "9 Migration configuration and operational readiness", 1)
heading(doc, "Client inputs required before build completion", 2)
inputs = [
    ["Organization structure", "Final property, cluster and legal-entity list with addresses, currency, financial year and active status"],
    ["Authority matrix", "Named roles, delegates, value thresholds, category and legal-entity variations, exception routes and segregation-of-duties decisions"],
    ["Budgets", "Approved CAPEX and operating budget workbooks, line references, period values, owners and opening commitments"],
    ["Items", "Clean item and service catalogue with SKU, UOM, category, tax, price, supplier and approval status"],
    ["Suppliers", "Legal and display names, contacts, categories, tax and bank data, property scope, status and approved compliance documents"],
    ["Opening transactions", "Approved list of open PRs, RFQs, POs, receipts and invoices to migrate, with balances and source references"],
    ["Templates", "Approved CAPEX, OpEx, repair, RFQ, award, rejection, PO, GRN, invoice and remittance layouts and branding"],
    ["Integration", "Books, email, WhatsApp, SSO and bank decisions plus credentials provided through approved secure channels"],
    ["Operations", "Service owners, support contacts, escalation routes, retention, backup, recovery and incident requirements"],
]
add_table(doc, ["Input", "Required content"], inputs, widths=[1.45, 5.15], font_size=8.5)

heading(doc, "Migration controls", 2)
for i, text in enumerate([
    "Provide approved import templates and data dictionaries before collecting final data.",
    "Run a profiling and cleansing pass that reports missing required values, invalid references, duplicates and inactive records.",
    "Perform at least one mock load and obtain client reconciliation of record counts, key financial totals and sampled documents.",
    "Freeze source data for final extraction or define a controlled delta process.",
    "Record every production import as a batch with source filename, checksum, operator, date, counts, rejects and reconciliation result.",
    "Retain rollback files and define how migrated transactions will be corrected without bypassing audit rules.",
], 1):
    add_number(doc, text, i)

heading(doc, "Operational deliverables", 2)
for text in [
    "Production architecture and environment inventory",
    "Configuration workbook and authority matrix",
    "Data dictionary and integration mapping",
    "Deployment, smoke-test and rollback runbook",
    "Backup and restoration procedure with test evidence",
    "Monitoring and incident-response procedure",
    "Administrator guide, end-user guide and supplier portal guide",
    "Role-based training material and attendance record",
    "Known limitations and approved workarounds",
    "Hypercare plan with issue ownership and exit criteria",
]:
    add_bullet(doc, text)

heading(doc, "10 Verification acceptance and traceability", 1)
paragraph(doc, "The revised SRS must include or reference a Requirements Traceability Matrix. Each Must requirement needs a test ID, owner, evidence and result. Demonstrations without recorded evidence should not count as acceptance.")

tests = [
    ["T01", "Authentication and access", "Unauthenticated requests fail; inactive users fail; permitted users see only authorized modules and properties."],
    ["T02", "Approval routes", "Every approved workflow variant is tested stage by stage, including vacancies, delegation, rejection, recall, resubmission and thresholds."],
    ["T03", "Budget control", "Within-budget, no-budget, exceeded-budget, concurrent submission, rejection release, invoice actualization and reversal reconcile correctly."],
    ["T04", "Document gates", "Every required document rule blocks progression when missing and allows progression when valid; waivers require authority and audit."],
    ["T05", "Supplier onboarding", "Incomplete registration is blocked; SCM and Finance actions route correctly; rejected suppliers cannot bid; active suppliers can access only their own data."],
    ["T06", "Sealed bids", "No unauthorized user can view bid content before opening; deadlines close submissions; revisions create versions; one supplier is awarded."],
    ["T07", "PO control", "PO value equals approved award or exception approval; issue version is retained; supplier accept or reject drives the correct state."],
    ["T08", "Receiving", "Partial and multiple receipts calculate cumulative accepted and rejected values; over-receipt is blocked; fulfillment waits for completion or short close."],
    ["T09", "Invoice match", "Line quantities, prices, taxes and totals are matched across partial GRNs and invoices; discrepancies and tolerances are explained."],
    ["T10", "Payments and credits", "Unresolved discrepancies block payment; authorized override works; credits reduce balance once; overpayment and double application fail."],
    ["T11", "Reports", "Report totals reconcile to source transactions and respect property, supplier and role scope."],
    ["T12", "Integrations", "Success, timeout, duplicate retry, rejected payload, expired credentials and replay paths are tested and reconciled."],
    ["T13", "Audit", "All material actions show actor, action, record, old and new state or detail, time and correlation reference; ordinary users cannot alter logs."],
    ["T14", "Security", "Authorization, property isolation, supplier isolation, file access, rate limiting, secret exposure and common web vulnerabilities are tested."],
    ["T15", "Recovery", "Backup restoration meets approved RPO and RTO and restored financial totals reconcile."],
    ["T16", "Performance and compatibility", "Approved load, browser and mobile test suites meet the targets in the final NFRs."],
]
add_table(doc, ["Test group", "Area", "Minimum evidence"], tests, widths=[0.65, 1.45, 4.5], font_size=8)

heading(doc, "Minimum end to end UAT scenarios", 2)
for text in [
    "Budgeted CapEx addition from request through approval, RFQ, sealed bids, award, PO, commissioning, asset creation, invoice match and payment.",
    "Non-budgeted request with approved reason and evidence through the required executive or Board route.",
    "Repair below and above each legal-entity threshold, proving the corporate-approval exception.",
    "Service or AMC request using SES rather than a goods receipt.",
    "Supplier invitation, registration, SCM review, Finance review, activation, bid and PO acknowledgement.",
    "Partial delivery with accepted and rejected quantities, later replacement delivery, partial invoice and final settlement.",
    "Invoice discrepancy blocked from payment, followed by authorized resolution or credit and successful payment.",
    "Supplier rejection of a PO followed by cancellation, negotiation or alternate award using the approved route.",
    "Property-level user attempts cross-property access through both the interface and direct API calls and is refused.",
    "Optional integration failure followed by retry and reconciliation without duplicating the business transaction.",
]:
    add_bullet(doc, text)

heading(doc, "Traceability matrix fields", 2)
add_table(doc, ["Field", "Required value"], [
    ["Requirement ID", "Stable unique identifier such as PR FR 001 or SEC NFR 004"],
    ["Source", "SRS section, approved workshop decision, law, policy or interface contract"],
    ["Priority and release", "Must, Should or Could plus target release"],
    ["Owner", "Business owner accountable for clarification and acceptance"],
    ["Design reference", "Screen, workflow, schema, interface or configuration reference"],
    ["Test reference", "Automated test, SIT case and UAT case"],
    ["Evidence", "Screenshot, exported record, audit event, report reconciliation or test output"],
    ["Result and approval", "Pass or fail, defect reference, approver and date"],
], widths=[1.45, 5.15], font_size=8.5)

heading(doc, "11 Definition of complete delivery", 1)
paragraph(doc, "ProcureFlow should be described as completely built for the approved Galle Face Hotel scope only when every condition below is satisfied.")
definition = [
    ["Requirements", "All P0 decisions are signed, the SRS has stable requirement IDs and every Must requirement has acceptance criteria."],
    ["Implementation", "All agreed workflows and validations are enforced by the server and represented correctly in the interface."],
    ["Data", "The production schema, reference data and migrated records are reconciled and approved by their business owners."],
    ["Security", "Identity, permissions, property isolation, supplier isolation, files, secrets and audit controls pass the approved tests."],
    ["Financial control", "Budget commitments, receipts, invoice matches, credits and payments reconcile without duplicate or unexplained balances."],
    ["Documents", "Required evidence gates, approved templates and issued-document versions operate for every in-scope category and legal entity."],
    ["Integrations", "Every in-scope integration passes success and failure tests and has an owner and reconciliation procedure."],
    ["Quality", "Automated tests, SIT and UAT pass; unresolved defects are explicitly accepted with owner, workaround and due date."],
    ["Operations", "Monitoring, backup, restore, deployment, rollback, support and incident procedures are tested and handed over."],
    ["Adoption", "Administrators, employees, approvers, procurement, stores, Finance and pilot suppliers complete role-based training."],
    ["Acceptance", "Procurement, Finance, IT or Security and the executive sponsor provide documented production approval."],
]
add_table(doc, ["Completion area", "Required evidence"], definition, widths=[1.45, 5.15], font_size=8.5)

heading(doc, "Items that must not be claimed before verification", 2)
for text in [
    "Exact compliance with the current SRS",
    "Automated bank payment execution",
    "Sealed-bid confidentiality",
    "Complete supplier onboarding and approval",
    "CAPEX workbook replication",
    "Mandatory evidence enforcement at every stage",
    "Full line-level three-way matching",
    "Transaction-safe budget concurrency",
    "Shared multi-tenant SaaS isolation",
    "Production availability or recovery targets",
]:
    add_bullet(doc, text)

heading(doc, "12 Instructions to the SRS author", 1)
for i, text in enumerate([
    "Create SRS Version 1.1 and record this revision brief in Document History.",
    "Run a decision workshop for items D01 to D10 and attach the signed decisions to the SRS baseline.",
    "Replace conflicting narrative workflows with authoritative matrices and state diagrams.",
    "Give every functional and nonfunctional requirement a stable ID, priority, owner and acceptance criteria.",
    "Add data definitions, state transitions, permission matrices, required-document matrices, report definitions and interface contracts as controlled appendices.",
    "Mark optional items and external dependencies clearly. Do not mix first-release commitments with possible future capabilities.",
    "Build a Requirements Traceability Matrix and obtain review from Procurement, Finance, IT or Security and the implementation lead.",
    "Issue the revised SRS for formal approval before final estimation and build commitment.",
], 1):
    add_number(doc, text, i)

heading(doc, "Required SRS appendices", 2)
appendices = [
    ["A", "Terminology and controlled values", "Roles, categories, budget states, transaction states, legal entities and properties"],
    ["B", "Authority matrix", "Workflow conditions, stages, roles, limits, exceptions, delegation and segregation of duties"],
    ["C", "Field dictionary", "Fields, types, lengths, required conditions, validation and source for each module"],
    ["D", "Document requirement matrix", "Required attachments by category, stage, legal entity, threshold and waiver authority"],
    ["E", "Status transition matrix", "Allowed actions, actors, prerequisites and results for every transaction state"],
    ["F", "Permission matrix", "Role and profile permissions plus property and supplier scope"],
    ["G", "Report catalogue", "Columns, formulas, filters, security, freshness and reconciliation"],
    ["H", "Integration specifications", "Authentication, mappings, triggers, retry, idempotency and support"],
    ["I", "Migration specification", "Sources, templates, validation, reconciliation, cutover and rollback"],
    ["J", "Requirements Traceability Matrix", "Requirement, design, development, test, evidence and approval"],
]
add_table(doc, ["Appendix", "Name", "Required content"], appendices, widths=[0.7, 2.0, 3.9], font_size=8.3)

add_page_break(doc)
heading(doc, "Appendix A Current ProcureFlow baseline and confirmed gaps", 1)
paragraph(doc, "This review compared the supplied SRS with the current ProcureFlow repository. The baseline contains substantial working functions, but the gaps below prevent an exact-compliance statement.")
baseline = [
    ["Property scoped procurement", "Supported baseline", "Property records and assignments exist and many queries are organization scoped.", "Complete legal-entity model, verify every endpoint and file download, and test property isolation."],
    ["Approval routes", "Enhancement required", "The code contains budgeted, non-budgeted and budget-exceeded route arrays.", "Replace fixed budget routes with the final category, entity and threshold rules and persist route versions."],
    ["Budget periods", "Supported baseline", "Budget and period records exist with committed and spent values.", "Add item-level CAPEX lines, quarterly classifications, ledger safety and exact SRS validation."],
    ["RFQ and awards", "Supported baseline", "Approved PR conversion, supplier invitations, bids and one award to PO exist.", "Add draft and publish separation, document gates, sealed opening, amendments, evaluations and controlled negotiation."],
    ["Supplier portal", "Supported baseline", "Invitations, hashed access, scoped RFQs, bid updates, POs, invoices and payments exist.", "Build registration, evidence, SCM and Finance review and approved activation."],
    ["Goods receipts", "Correction required", "Quantity fields and over-receipt control exist.", "Require accepted plus rejected equals received and do not mark a partial receipt Fulfilled."],
    ["Invoice matching", "Enhancement required", "Current logic compares PO totals and accepted GRN quantities.", "Capture invoice lines and match quantities, prices and taxes across partial receipts and invoices."],
    ["Payments", "Correction required", "Partial and full recording and overpayment control exist.", "Block unresolved discrepancies, implement override authorization and apply supplier credits."],
    ["Attachments", "Supported baseline", "Attachment upload and secure storage paths exist.", "Add stage-specific required slots, validation, waivers, retention and malware decision."],
    ["Audit and support", "Supported baseline", "Audit records and constrained support-session test assets exist.", "Complete append-only guarantees, sensitive events, production administration and SRS architecture alignment."],
    ["Books integration", "Supported baseline", "Configuration, OAuth, test, disconnect and sync endpoints exist.", "Finalize field mapping, source-of-truth, retries, idempotency and financial reconciliation."],
    ["Verification", "Partial evidence", "Existing local workflow and API suites passed 66 checks during this review.", "Add tests for every revised SRS rule and run them against a controlled Catalyst environment and production-like data."],
]
add_status_table(doc, baseline)

heading(doc, "Material code behaviors observed", 2)
for text in [
    "Purchase-request submission currently checks an active department and property budget and can reject an over-budget amount before the budget-exceeded route is used. The final SRS and implementation must agree on whether excess requests are blocked or escalated.",
    "The receipt endpoint currently updates the PO to Fulfilled after saving a receipt. This must be changed for partial deliveries.",
    "The receipt validation currently rejects accepted plus rejected greater than received; the SRS requires equality.",
    "The invoice matcher operates mainly at PO total and accepted-quantity level because invoice lines are not represented as a full matching structure.",
    "The payment endpoint blocks overpayment but does not currently enforce the required discrepancy block and authorized override.",
    "The repository deployment guide explicitly describes one Catalyst project for one hotel group, while the SRS describes a shared multi-tenant product.",
]:
    add_bullet(doc, text)

heading(doc, "Appendix B Recommended requirement wording patterns", 1)
paragraph(doc, "Use the following pattern to rewrite each requirement. Replace bracketed values with approved business decisions.")

heading(doc, "Functional requirement pattern", 2)
paragraph(doc, "[Requirement ID] When [trigger] occurs, the system shall allow [authorized role] to [action] only if [preconditions]. The system shall validate [rules], store [data], change the status from [old state] to [new state], write an audit event containing [fields], notify [recipient] through [channel], and return [defined error behavior] when validation fails.")

heading(doc, "Acceptance criterion pattern", 2)
paragraph(doc, "Given [starting data and actor], when [action] is performed, then [observable result] shall occur and [prohibited result] shall not occur. Evidence shall include [record, audit event, rendered document, report reconciliation or integration response].")

heading(doc, "Example receipt requirement", 2)
paragraph(doc, "GRN FR 007 When an authorized receiving user records a receipt against an Issued or Partially Received purchase order, the system shall require Received Quantity to equal Accepted Quantity plus Rejected Quantity for every line. The cumulative Received Quantity shall not exceed Ordered Quantity unless an approved over-receipt tolerance applies. The purchase order shall remain Partially Received until cumulative Accepted Quantity equals Ordered Quantity or an authorized short-close decision is recorded. The system shall audit the quantities, actor, evidence and resulting status.")

heading(doc, "Example discrepancy payment requirement", 2)
paragraph(doc, "PAY FR 004 The system shall reject a payment against an invoice in Discrepancy status. A Finance user with the mismatch-override permission may change the invoice to Override Approved only after recording a reason and attaching the required evidence. The audit log shall identify the Finance user, discrepancy values, reason, evidence references and approval time.")

heading(doc, "Source basis", 2)
for text in [
    "GalleFace Hotel SRS.pdf, 79 pages, supplied for review on 17 September 2026.",
    "Current ProcureFlow repository in C:\\Users\\MK\\Documents\\Procurement, including Catalyst functions, web client, deployment guidance and verification assets.",
    "Zoho Catalyst platform capabilities used by the existing solution include hosted authentication, Advanced I/O Functions, Data Store and Stratus object storage.",
]:
    add_bullet(doc, text)

# Footer with page numbers
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Galle Face Hotel ProcureFlow SRS Revision and Development Requirements   |   ", size=8, color=MIDGRAY)
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fldChar1)
    r._r.append(instrText)
    r._r.append(fldChar2)
    set_run_font(r, size=8, color=MIDGRAY)

# Metadata
doc.core_properties.title = "Galle Face Hotel ProcureFlow SRS Revision and Development Requirements"
doc.core_properties.subject = "SRS revision instructions, complete development scope and acceptance baseline"
doc.core_properties.author = "ProcureFlow Review"
doc.core_properties.keywords = "ProcureFlow, Galle Face Hotel, SRS, Zoho Catalyst, requirements, development"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
